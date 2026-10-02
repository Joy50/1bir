from datetime import date, datetime, timedelta
from pathlib import Path
import re

from django.core.exceptions import ValidationError
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from common.models import MobileNumber, Organization, Person, Rank
from common.scoping import get_or_create_hq_company, get_unit


PERSONNEL_SHEETS = (
    "OFFRS",
    "MWO",
    "SWO",
    "WO",
    "SGT",
    "CPL",
    "LCPL",
    "SNK",
    "Clk",
    "CK",
    "NCE",
    "NC (U)",
)
COMPANIES = ("A Company", "B Company", "C Company", "D Company", "HQ Company")
ROLL_DATE = date(2025, 7, 1)
RANK_ALIASES = {
    "lt col": "Lt. Col",
    "lt. col": "Lt. Col",
    "ltcol": "Lt. Col",
    "maj": "Maj",
    "capt": "Capt",
    "lt": "Lt",
    "mwo": "MWO",
    "swo": "SWO",
    "wo": "WO",
    "h capt": "H Capt",
    "h lt": "H Lt",
    "sgt": "Sgt",
    "cpl": "Cpl",
    "lcpl": "Lcpl",
    "lcpl/m": "Lcpl",
    "snk": "Snk",
    "snk /m": "Snk",
    "snk/m": "Snk",
    "nce": "NCE",
    "nc (u)": "NC (U)",
    "nc(u)": "NC (U)",
}
OFFICER_RANKS = {"Lt. Col", "Maj", "Capt", "Lt", "Col"}
JCO_RANKS = {"MWO", "SWO", "WO", "H Capt", "H Lt"}


def _norm_header(value):
    text = " ".join(str(value or "").lower().replace("\n", " ").split())
    return text.replace(".", "")


def header_key(label):
    header = _norm_header(label)
    if not header:
        return None
    if header.startswith("army no"):
        return "army_number"
    if header in {"rank", "rk"}:
        return "rank"
    if header == "name":
        return "name"
    if header == "unit":
        return "unit"
    if header == "coy":
        return "coy"
    if header == "dob":
        return "dob"
    if header in {"doe", "doc"}:
        return "doe"
    if header == "lc":
        return "batch"
    if header.startswith("ai "):
        return "al1_13"
    if header == "dor":
        return "dor"
    if header.startswith("discipline"):
        return "discipline"
    if header.startswith("punishment") and "dt" not in header:
        return "punishment"
    if header.startswith("msn"):
        return "mission"
    if "present age" in header:
        return "age"
    if "present svc" in header:
        return "svc_years"
    if header.startswith("nid"):
        return "nid"
    if header.startswith("ime"):
        return "imei"
    if header.startswith("mobile no-1") or header.startswith("mobile no 1"):
        return "mobile"
    return None


def clean_text(value, max_length=None):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    if text.lower() in {"", "-", "n/a", "na", "none", "nil"}:
        return ""
    if max_length:
        return text[:max_length]
    return text


def army_number(value):
    text = clean_text(value, 20)
    if not text or text.lower() in {"army no", "ser"}:
        return ""
    if re.fullmatch(r"\d+\.0", text):
        text = text[:-2]
    if re.fullmatch(r"\d+\.0+", text):
        text = text.split(".")[0]
    return text[:20]


def parse_date(value):
    if value in (None, "", "-"):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, int) and 1900 < value < 2100:
        return None
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d-%m-%y", "%d/%m/%y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def parse_int(value, minimum=None, maximum=None):
    try:
        number = int(float(value))
    except (TypeError, ValueError):
        return None
    if minimum is not None and number < minimum:
        return None
    if maximum is not None and number > maximum:
        return None
    return number


def infer_date(years):
    if years is None:
        return None
    try:
        return ROLL_DATE.replace(year=ROLL_DATE.year - years)
    except ValueError:
        return date(ROLL_DATE.year - years, 7, 1)


def yes_no(value):
    text = clean_text(value).lower()
    if text in {"yes", "y"}:
        return "Yes"
    if text in {"no", "n"}:
        return "No"
    return None


def parse_mission(value):
    text = clean_text(value).lower()
    if text in {"done", "yes", "y"}:
        return True
    if text in {"not done", "no", "n"}:
        return False
    return None


def normalize_rank(value):
    text = clean_text(value)
    if not text:
        return ""
    key = " ".join(text.lower().replace(".", " ").split())
    return RANK_ALIASES.get(key, text)


def rank_category(name):
    if name in OFFICER_RANKS:
        return Rank.CATEGORY_OFFICER
    if name in JCO_RANKS:
        return Rank.CATEGORY_JCO
    return Rank.CATEGORY_OR


def map_company(coy, unit):
    raw = clean_text(coy) or clean_text(unit)
    key = (
        raw.lower()
        .replace("company", " ")
        .replace("coy", " ")
        .replace(".", " ")
        .strip()
    )
    if not key or parse_date(coy):
        raw = clean_text(unit)
        key = raw.lower().strip()
    if key.startswith("a") and "brig" not in key:
        return "A Company"
    if key.startswith("b") and "bde" not in key and "bir" not in key:
        return "B Company"
    if key.startswith("c") and "cs-" not in key:
        return "C Company"
    if key.startswith("d"):
        return "D Company"
    if key.startswith("hq") or key in {"h q", "headquarter", "headquarters"}:
        return "HQ Company"
    return "HQ Company"


def format_mobile(value):
    text = clean_text(value)
    if not text:
        return ""
    digits = re.sub(r"\D", "", text)
    if len(digits) == 10 and digits.startswith("1"):
        digits = "0" + digits
    if 10 <= len(digits) <= 15:
        return digits
    if 8 <= len(text) <= 15:
        return text
    return ""


class Command(BaseCommand):
    help = "Import soldiers from the 1 BIR seniority roll workbook."

    def add_arguments(self, parser):
        parser.add_argument(
            "path",
            nargs="?",
            default=str(
                Path.home()
                / "Downloads"
                / "Sr Roll-JCO & ORs (GD)-1 BIR-01-7-2025.xlsx"
            ),
        )

    def handle(self, *args, **options):
        path = Path(options["path"])
        if not path.exists():
            raise CommandError(f"File not found: {path}")
        try:
            import openpyxl
        except ImportError as exc:
            raise CommandError("Install openpyxl first: pip install openpyxl") from exc

        companies = self._ensure_organizations()
        workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
        created = updated = skipped = 0
        skip_reasons = []
        try:
            for sheet_name in PERSONNEL_SHEETS:
                if sheet_name not in workbook.sheetnames:
                    skip_reasons.append(f"{sheet_name}: sheet missing")
                    continue
                sheet_created, sheet_updated, sheet_skipped, reasons = self._import_sheet(
                    workbook[sheet_name],
                    sheet_name,
                    companies,
                )
                created += sheet_created
                updated += sheet_updated
                skipped += sheet_skipped
                skip_reasons.extend(reasons)
                self.stdout.write(
                    f"{sheet_name}: +{sheet_created} new, {sheet_updated} updated, "
                    f"{sheet_skipped} skipped"
                )
        finally:
            workbook.close()

        self.stdout.write(
            self.style.SUCCESS(
                f"Done. {created} created, {updated} updated, {skipped} skipped. "
                f"Total soldiers now: {Person.objects.count()}"
            )
        )
        for reason in skip_reasons[:25]:
            self.stdout.write(f"  skip: {reason}")
        if len(skip_reasons) > 25:
            self.stdout.write(f"  ... {len(skip_reasons) - 25} more skips")

    def _ensure_organizations(self):
        unit = get_unit()
        if unit is None:
            unit = Organization.objects.create(
                organization_name="1 BIR",
                unit_kind=Organization.KIND_UNIT,
            )
        elif unit.organization_name.lower() == "1 bir":
            unit.organization_name = "1 BIR"
            unit.save(update_fields=["organization_name"])
        companies = {}
        for name in COMPANIES:
            company, _created = Organization.objects.get_or_create(
                organization_name=name,
                parent_organization=unit,
                defaults={"unit_kind": Organization.KIND_COMPANY},
            )
            if company.unit_kind != Organization.KIND_COMPANY:
                company.unit_kind = Organization.KIND_COMPANY
                company.save(update_fields=["unit_kind"])
            companies[name] = company
        get_or_create_hq_company(unit)
        return companies

    def _import_sheet(self, worksheet, sheet_name, companies):
        mapping = None
        created = updated = skipped = 0
        reasons = []
        for row in worksheet.iter_rows(values_only=True):
            if mapping is None:
                if any(
                    str(cell or "").strip().lower().startswith("army no")
                    for cell in row[:8]
                ):
                    mapping = {}
                    for index, cell in enumerate(row):
                        key = header_key(cell)
                        if key and key not in mapping:
                            mapping[key] = index
                continue
            if "army_number" not in mapping or "name" not in mapping:
                reasons.append(f"{sheet_name}: no Army No/Name columns")
                break
            record = {
                key: row[index] if index < len(row) else None
                for key, index in mapping.items()
            }
            number = army_number(record.get("army_number"))
            name = clean_text(record.get("name"), 100)
            if not number and not name:
                continue
            if not number or not name:
                skipped += 1
                reasons.append(f"{sheet_name}: incomplete row {number or name}")
                continue
            rank_name = normalize_rank(record.get("rank"))
            if not rank_name:
                skipped += 1
                reasons.append(f"{number}: missing rank")
                continue
            dob = parse_date(record.get("dob")) or infer_date(
                parse_int(record.get("age"), 16, 70)
            )
            doe = parse_date(record.get("doe")) or infer_date(
                parse_int(record.get("svc_years"), 0, 45)
            )
            if dob and not doe:
                try:
                    doe = dob.replace(year=dob.year + 18)
                except ValueError:
                    doe = date(dob.year + 18, 3, 1)
            if doe and not dob:
                try:
                    dob = doe.replace(year=doe.year - 18)
                except ValueError:
                    dob = date(doe.year - 18, 3, 1)
            if not dob or not doe:
                dob = dob or date(1995, 1, 1)
                doe = doe or date(2013, 1, 2)
            if doe <= dob:
                doe = dob + timedelta(days=1)
            company = companies[map_company(record.get("coy"), record.get("unit"))]
            rank = self._get_rank(rank_name)
            defaults = {
                "name": name,
                "rank": rank,
                "organization": company,
                "dob": dob,
                "doe": doe,
                "batch": clean_text(record.get("batch"), 30),
                "al1_13": yes_no(record.get("al1_13")),
                "dor": parse_date(record.get("dor")),
                "discipline": clean_text(record.get("discipline")) or None,
                "punishment": clean_text(record.get("punishment")) or None,
                "mission": parse_mission(record.get("mission")) or False,
                "nid_number": clean_text(record.get("nid"), 30),
                "phone_imei": clean_text(record.get("imei"), 50),
            }
            try:
                with transaction.atomic():
                    person, was_created = Person.objects.update_or_create(
                        army_number=number,
                        defaults=defaults,
                    )
                    person.full_clean()
                    person.save()
                    self._save_mobile(person, record.get("mobile"))
            except (ValidationError, ValueError) as exc:
                skipped += 1
                reasons.append(f"{number}: {exc}")
                continue
            if was_created:
                created += 1
            else:
                updated += 1
        return created, updated, skipped, reasons

    def _get_rank(self, name):
        rank, created = Rank.objects.get_or_create(
            rank_name=name,
            defaults={"category": rank_category(name)},
        )
        if not created and rank.category != rank_category(name) and name in OFFICER_RANKS | JCO_RANKS:
            rank.category = rank_category(name)
            rank.save(update_fields=["category"])
        return rank

    def _save_mobile(self, person, value):
        number = format_mobile(value)
        if not number:
            return
        MobileNumber.objects.get_or_create(
            person=person,
            mobile_number=number,
            defaults={"type_of_number": MobileNumber.TYPE_PERSONAL},
        )
